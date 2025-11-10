from docx import Document

doc = Document('/vercel/sandbox/uploads/Nallagatla Vamshi 35.docx')

for para in doc.paragraphs:
    print(para.text)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                print(cell.text)
