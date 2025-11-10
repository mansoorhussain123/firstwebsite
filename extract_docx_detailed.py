from docx import Document
import json

doc = Document('/vercel/sandbox/uploads/Nallagatla Vamshi 35.docx')

print("=== PARAGRAPHS ===")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"Para {i}: {para.text}")
    for run in para.runs:
        if run.text.strip():
            print(f"  Run: {run.text}")

print("\n=== TABLES ===")
for table_idx, table in enumerate(doc.tables):
    print(f"\nTable {table_idx}:")
    for row_idx, row in enumerate(table.rows):
        row_data = []
        for cell in row.cells:
            row_data.append(cell.text.strip())
        if any(row_data):
            print(f"  Row {row_idx}: {row_data}")

print("\n=== SECTIONS ===")
for section in doc.sections:
    print(f"Section: {section}")

print("\n=== CORE PROPERTIES ===")
core_props = doc.core_properties
print(f"Title: {core_props.title}")
print(f"Author: {core_props.author}")
print(f"Subject: {core_props.subject}")
