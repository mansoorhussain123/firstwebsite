from docx import Document
import zipfile

# Check if it's a valid DOCX
try:
    with zipfile.ZipFile('uploads/Nallagatla Vamshi 35.docx', 'r') as zip_ref:
        print("Files in DOCX:")
        for name in zip_ref.namelist()[:10]:
            print(f"  {name}")
    
    print("\n--- Document Content ---")
    doc = Document('uploads/Nallagatla Vamshi 35.docx')
    
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Sections: {len(doc.sections)}")
    
    # Try to extract all text
    full_text = []
    for para in doc.paragraphs:
        if para.text:
            full_text.append(para.text)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text:
                        full_text.append(para.text)
    
    if full_text:
        print("\n--- Resume Text ---")
        print('\n'.join(full_text))
    else:
        print("No text found in document")
        
except Exception as e:
    print(f"Error: {e}")
